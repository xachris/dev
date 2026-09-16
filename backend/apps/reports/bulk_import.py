import csv
import io
from dataclasses import dataclass

from django.core.exceptions import ValidationError
from openpyxl import load_workbook


MAX_IMPORT_ROWS = 500

HEADER_ALIASES = {
    "student_number": "student_number",
    "学号": "student_number",
    "student_name": "student_name",
    "姓名": "student_name",
    "student_feedback": "student_feedback",
    "学生反馈": "student_feedback",
    "guardian_message": "guardian_message",
    "家长信息": "guardian_message",
    "staff_note": "staff_note",
    "内部备注": "staff_note",
    "action": "action",
    "动作": "action",
}

REQUIRED_HEADERS = {"student_number", "student_feedback"}


@dataclass
class ImportRow:
    row_number: int
    student_number: str
    student_name: str
    student_feedback: str
    guardian_message: str
    staff_note: str
    action: str

    def as_dict(self):
        return {
            "row_number": self.row_number,
            "student_number": self.student_number,
            "student_name": self.student_name,
            "student_feedback": self.student_feedback,
            "guardian_message": self.guardian_message,
            "staff_note": self.staff_note,
            "action": self.action,
        }


def _normalise_headers(headers):
    normalised = []
    for header in headers:
        key = str(header or "").strip()
        normalised.append(HEADER_ALIASES.get(key, key))
    return normalised


def _coerce(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_from_csv(upload):
    raw = upload.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationError("CSV 必须使用 UTF-8 编码。") from exc
    reader = csv.reader(io.StringIO(text))
    return list(reader)


def _rows_from_xlsx(upload):
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("无法读取 Excel 文件，请确认文件为有效的 .xlsx。") from exc
    worksheet = workbook.active
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def parse_import_upload(upload):
    filename = (upload.name or "").lower()
    if filename.endswith(".csv"):
        raw_rows = _rows_from_csv(upload)
    elif filename.endswith(".xlsx"):
        raw_rows = _rows_from_xlsx(upload)
    else:
        raise ValidationError("仅支持 .csv 或 .xlsx 文件。")

    if not raw_rows:
        raise ValidationError("文件为空。")

    headers = _normalise_headers(raw_rows[0])
    if len(set(headers)) != len(headers):
        raise ValidationError("表头存在重复字段。")

    missing = REQUIRED_HEADERS - set(headers)
    if missing:
        raise ValidationError(f"缺少必需字段：{', '.join(sorted(missing))}")

    rows = []
    duplicate_student_numbers = set()
    seen_student_numbers = set()

    for index, raw_row in enumerate(raw_rows[1:], start=2):
        if len(rows) >= MAX_IMPORT_ROWS:
            raise ValidationError(f"单次最多导入 {MAX_IMPORT_ROWS} 行。")

        padded = list(raw_row) + [""] * max(0, len(headers) - len(raw_row))
        data = {headers[i]: _coerce(padded[i]) for i in range(len(headers))}

        if not any(data.values()):
            continue

        student_number = data.get("student_number", "")
        if not student_number:
            raise ValidationError(f"第 {index} 行缺少学号。")
        if student_number in seen_student_numbers:
            duplicate_student_numbers.add(student_number)
        seen_student_numbers.add(student_number)

        action = data.get("action", "DRAFT").strip().upper() or "DRAFT"
        if action not in {"DRAFT", "SUBMIT"}:
            raise ValidationError(f"第 {index} 行 action 只能是 DRAFT 或 SUBMIT。")

        rows.append(
            ImportRow(
                row_number=index,
                student_number=student_number,
                student_name=data.get("student_name", ""),
                student_feedback=data.get("student_feedback", ""),
                guardian_message=data.get("guardian_message", ""),
                staff_note=data.get("staff_note", ""),
                action=action,
            )
        )

    if duplicate_student_numbers:
        duplicates = ", ".join(sorted(duplicate_student_numbers))
        raise ValidationError(f"文件内学号重复：{duplicates}")
    if not rows:
        raise ValidationError("文件中没有可导入的数据行。")
    return rows
